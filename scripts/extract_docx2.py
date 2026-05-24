import sys
from docx import Document

doc = Document(sys.argv[1])
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else "Normal"
    text = p.text.strip()
    if text:
        print(f"[{i}] ({style}) {text}")

print()
print(f"TABLES: {len(doc.tables)}")
for ti, t in enumerate(doc.tables):
    print(f"--- Table {ti} ({len(t.rows)}x{len(t.columns)}) ---")
    for ri, row in enumerate(t.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"  R{ri}: {cells}")
