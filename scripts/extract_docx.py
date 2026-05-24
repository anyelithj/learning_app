import sys
from docx import Document

doc = Document(sys.argv[1])

print("=" * 80)
print("PARAGRAPHS / HEADINGS")
print("=" * 80)
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else "Normal"
    text = p.text.strip()
    if text or style.startswith("Heading"):
        print(f"[{i}] ({style}) {text[:200]}")

print()
print("=" * 80)
print(f"TABLES: {len(doc.tables)}")
print("=" * 80)
for ti, t in enumerate(doc.tables):
    print(f"--- Table {ti} ({len(t.rows)}x{len(t.columns)}) ---")
    for ri, row in enumerate(t.rows[:5]):
        cells = [c.text.strip()[:60] for c in row.cells]
        print(f"  R{ri}: {cells}")
